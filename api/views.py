from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.models import User
from .models import Prediction, Sheet
from django.http import JsonResponse
from django.middleware.csrf import get_token
from django.views.decorators.csrf import ensure_csrf_cookie
from django.http import HttpResponse
from openpyxl import load_workbook
from io import BytesIO
import json

with open('api/codes.json', 'r', encoding='utf-8') as f:
    all_codes: set = set(json.load(f))
with open('api/matchmap.json', 'r', encoding='utf-8') as f:
    excel_map: dict = json.load(f)

@ensure_csrf_cookie
def csrf(request):
    return JsonResponse({"csrfToken": get_token(request)})

def me(request):
    if request.user.is_authenticated:
        sheet = Sheet.objects.filter(user=request.user).first()
        return JsonResponse({
            "authenticated": True,
            "username": request.user.username,
            "email": request.user.email,
            "firstSheetId": sheet.id if sheet else None
        })
    return JsonResponse({"authenticated": False})

def register_user(request):
    if request.method != "POST":
        return JsonResponse({'success': False, 'message': 'Invalid request'})

    data = json.loads(request.body)
    username = data.get("username")
    email = data.get("email")
    password = data.get("password")

    # 1. Username exists?
    if User.objects.filter(username=username).exists():
        return JsonResponse({'success': False, 'message': 'Username not valid.'})

    # 2. Email exists?
    if User.objects.filter(email=email).exists():
        return JsonResponse({'success': False, 'message': 'Email not valid.'})

    try:
        # Create user
        user = User.objects.create_user(
            username=username,
            email=email,
            password=password
        )

        # Authenticate newly created user
        authenticated_user = authenticate(
            request,
            username=username,
            password=password
        )

        # Log them in
        if authenticated_user is not None:
            login(request, authenticated_user)

        # ⭐ Automatically create their first sheet
        Sheet.objects.create(
            user=authenticated_user,
            code=None  # or simply omit since null=True, blank=True
        )

        return JsonResponse({'success': True})

    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})

def login_user(request):
    if request.method != 'POST':
        return JsonResponse({"success": False, "message": "Invalid request"}, status=400)

    try:
        data = json.loads(request.body)
    except:
        return JsonResponse({"success": False, "message": "Invalid JSON"}, status=400)

    email = data.get("email")
    password = data.get("password")

    # Look up user by email
    user_obj = User.objects.filter(email=email).first()
    if not user_obj:
        return JsonResponse({"success": False, "message": "Invalid email or password"}, status=400)

    # Authenticate using username
    user = authenticate(request, username=user_obj.username, password=password)

    if user is None:
        return JsonResponse({"success": False, "message": "Invalid email or password"}, status=400)

    login(request, user)
    return JsonResponse({"success": True})

def logout_user(request):
    if request.method == "POST":
        logout(request)
        return JsonResponse({"success": True})
    return JsonResponse({"success": False, "message": "Invalid request"}, status=400)

def predict(request):
    if not request.user.is_authenticated:
        return JsonResponse({"success": False, "message": "Not authenticated"}, status=403)

    if request.method == "GET":
        sheet_id = request.GET.get("sheet_id")
        if not sheet_id:
            return JsonResponse({"success": False, "message": "sheet_id required"}, status=400)

        sheet = Sheet.objects.filter(id=sheet_id, user=request.user).first()
        if not sheet:
            return JsonResponse({"success": False, "message": "Sheet not found"}, status=404)

        preds = Prediction.objects.filter(sheet=sheet)

        predictions = {
            str(p.match_id): {
                "home_score": p.home_score,
                "away_score": p.away_score
            }
            for p in preds
        }

        return JsonResponse({
            "success": True,
            "submitted": sheet.code != None,
            "predictions": predictions
        })

    elif request.method == "POST":
        # SCORES LOCKED.
        return JsonResponse({"success": False, "message": "Scores locked."}, status=403)

        body = json.loads(request.body)
        sheet_id = body.get("sheet_id")
        predictions = body.get("predictions")

        if not sheet_id or not predictions:
            return JsonResponse({"success": False, "message": "sheet_id and predictions required"}, status=400)

        sheet = Sheet.objects.filter(id=sheet_id, user=request.user).first()
        if not sheet:
            return JsonResponse({"success": False, "message": "Sheet not found"}, status=404)

        for match_id, pred in predictions.items():
            Prediction.objects.update_or_create(
                sheet=sheet,
                match_id=int(match_id),
                defaults={
                    "home_score": pred["home_score"],
                    "away_score": pred["away_score"]
                }
            )

        return JsonResponse({"success": True})

    return JsonResponse({"success": False, "message": "Invalid request"}, status=400)

def my_sheets(request):
    if not request.user.is_authenticated:
        return JsonResponse({"success": False, "message": "Not authenticated"}, status=403)

    sheets = Sheet.objects.filter(user=request.user)

    data = [
        {
            "id": s.id,
            "owner": request.user.username,
            "rank": None,      # You can compute this later
            "submitted": s.code != None
        }
        for s in sheets
    ]

    return JsonResponse({"success": True, "sheets": data})

def other_sheets(request):
    authed_user = None
    if request.user.is_authenticated:
        authed_user = request.user

    # All sheets that:
    # 1. Are NOT owned by the current user
    # 2. Have a code (meaning they are submitted)
    sheets = Sheet.objects.filter(code__isnull=False).exclude(user=authed_user)

    data = [
        {
            "id": s.id,
            "owner": s.user.username,
            "rank": None,          # You can compute this later
            "submitted": True      # Always true because code != None
        }
        for s in sheets
    ]

    return JsonResponse({"success": True, "sheets": data})

def create_sheet(request):
    # SCORES LOCKED.
    return JsonResponse({"success": False, "message": "Scores locked."}, status=403)

    if request.method != "POST":
        return JsonResponse({"success": False, "message": "Invalid request"}, status=400)

    if not request.user.is_authenticated:
        return JsonResponse({"success": False, "message": "Not authenticated"}, status=403)

    sheet = Sheet.objects.create(user=request.user)

    return JsonResponse({
        "success": True,
        "sheet": {
            "id": sheet.id,
            "owner": request.user.username,
            "rank": None,
            "submitted": False
        }
    })

def submitted_count(request):
    count = Sheet.objects.exclude(code=None).count()

    return JsonResponse({
        "success": True,
        "count": count
    })

def submit(request):
    # SCORES LOCKED.
    return JsonResponse({"success": False, "message": "Scores locked."}, status=403)

    if request.method != "POST":
        return JsonResponse({"success": False, "message": "Invalid request"}, status=400)

    if not request.user.is_authenticated:
        return JsonResponse({"success": False, "message": "Not authenticated"}, status=403)

    # Parse JSON body
    try:
        body = json.loads(request.body)
    except:
        return JsonResponse({"success": False, "message": "Invalid JSON"}, status=400)

    code = body.get("code")
    sheet_id = body.get("sheet_id")

    if not code or not sheet_id:
        return JsonResponse({"success": False, "message": "code and sheet_id required"}, status=400)

    code = str(code).upper()

    # Validate code
    if code not in all_codes:
        return JsonResponse({"success": False, "message": "Bad code"}, status=403)

    # Check if code already used
    if Sheet.objects.filter(code=code).exists():
        return JsonResponse({"success": False, "message": "Code already used"}, status=403)

    # Get the sheet
    sheet = Sheet.objects.filter(id=sheet_id, user=request.user).first()
    if not sheet:
        return JsonResponse({"success": False, "message": "Sheet not found"}, status=404)

    # Check if already submitted
    if sheet.code is not None:
        return JsonResponse({"success": False, "message": "Sheet already submitted"}, status=403)

    # Submit the sheet
    sheet.code = code
    sheet.save()

    return JsonResponse({"success": True})

def unsubmit(request):
    # SCORES LOCKED.
    return JsonResponse({"success": False, "message": "Scores locked."}, status=403)

    if request.method != "POST":
        return JsonResponse({"success": False, "message": "Invalid request"}, status=400)

    if not request.user.is_authenticated:
        return JsonResponse({"success": False, "message": "Not authenticated"}, status=403)

    # Parse JSON body
    try:
        body = json.loads(request.body)
    except:
        return JsonResponse({"success": False, "message": "Invalid JSON"}, status=400)

    sheet_id = body.get("sheet_id")
    if not sheet_id:
        return JsonResponse({"success": False, "message": "sheet_id required"}, status=400)

    # Get the sheet
    sheet = Sheet.objects.filter(id=sheet_id, user=request.user).first()
    if not sheet:
        return JsonResponse({"success": False, "message": "Sheet not found"}, status=404)

    # Must already be submitted
    if sheet.code is None:
        return JsonResponse({"success": False, "message": "Sheet is not submitted"}, status=403)

    # Unsubmit the sheet
    sheet.code = None
    sheet.save()

    return JsonResponse({"success": True})

def export_sheet(request, sheet_id):
    sheet = Sheet.objects.get(id=sheet_id)

    # Build dict: match_id -> Prediction object
    preds = {p.match_id: p for p in Prediction.objects.filter(sheet=sheet)}

    wb = load_workbook("api/Polla2026.xlsx")
    ws = wb.active

    ws['U2'] = str(sheet.user.username)

    for match_id in range(1, 73):

        # Skip if match not in map
        if str(match_id) not in excel_map:
            continue

        pred = preds.get(match_id)
        if not pred:
            continue

        home_pred = pred.home_score
        away_pred = pred.away_score

        if home_pred is not None:
            ws[excel_map[str(match_id)][0]] = home_pred
        
        if away_pred is not None:
            ws[excel_map[str(match_id)][1]] = away_pred

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{sheet.user.get_username()}_{sheet_id}.xlsx"'
    return response